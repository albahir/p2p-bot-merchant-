"""Tests unitarios para el motor fiscal (FIFO vs Costo Promedio y generación de informe Excel)."""

import os
import openpyxl
from datetime import datetime, timezone
from src.data.models import Order
from src.reporting.tax_engine import FIFOCostMethod, WeightedAverageCostMethod
from src.reporting.excel_reporter import ExcelReporter

def test_fifo_vs_average_cost_calculations():
    orders = [
        Order(order_number="BUY_1", trade_type="BUY", asset="USDT", fiat="VES", fiat_amount=3500.0, crypto_amount=100.0, unit_price=35.0, fee_fiat=10.0, bcv_rate=35.0, order_status="COMPLETED", created_at=datetime.now(timezone.utc)),
        Order(order_number="BUY_2", trade_type="BUY", asset="USDT", fiat="VES", fiat_amount=4000.0, crypto_amount=100.0, unit_price=40.0, fee_fiat=10.0, bcv_rate=40.0, order_status="COMPLETED", created_at=datetime.now(timezone.utc)),
        Order(order_number="SELL_1", trade_type="SELL", asset="USDT", fiat="VES", fiat_amount=6750.0, crypto_amount=150.0, unit_price=45.0, fee_fiat=15.0, bcv_rate=40.0, order_status="COMPLETED", created_at=datetime.now(timezone.utc)),
    ]

    fifo = FIFOCostMethod()
    fifo_results = fifo.process_orders(orders)
    assert len(fifo_results) == 1
    assert fifo_results[0].acquisition_cost_fiat == 5515.0
    assert fifo_results[0].net_profit_fiat == 1220.0

    avg = WeightedAverageCostMethod()
    avg_results = avg.process_orders(orders)
    assert len(avg_results) == 1
    assert avg_results[0].acquisition_cost_fiat == 5640.0
    assert avg_results[0].net_profit_fiat == 1095.0

    assert fifo_results[0].net_profit_fiat != avg_results[0].net_profit_fiat

def test_excel_report_generation(tmp_path):
    orders = [
        Order(order_number="BUY_1", trade_type="BUY", asset="USDT", fiat="VES", fiat_amount=3500.0, crypto_amount=100.0, unit_price=35.0, fee_fiat=0.0, bcv_rate=35.0, order_status="COMPLETED", created_at=datetime.now(timezone.utc)),
        Order(order_number="SELL_1", trade_type="SELL", asset="USDT", fiat="VES", fiat_amount=3700.0, crypto_amount=100.0, unit_price=37.0, fee_fiat=0.0, bcv_rate=35.0, order_status="COMPLETED", created_at=datetime.now(timezone.utc)),
    ]
    fifo = FIFOCostMethod()
    results = fifo.process_orders(orders)

    output_file = str(tmp_path / "test_tax_report.xlsx")
    ExcelReporter.generate_tax_report(results, output_file)

    assert os.path.exists(output_file)
    wb = openpyxl.load_workbook(output_file)
    assert "Detalle" in wb.sheetnames
    assert "Resumen" in wb.sheetnames
